"""Check deliverable structure and numbers without claiming primary-data compliance."""
import csv
import hashlib
import json
from pathlib import Path
import zipfile
from docx import Document
from pptx import Presentation
import nbformat

ROOT=Path(__file__).resolve().parents[1]


def main():
    folder=ROOT/'submission'
    rows=list(csv.DictReader((folder/'dataset/labels.csv').open()))
    assert len({r['sha256'] for r in rows})==len(rows)
    assert len({r['file'] for r in rows})==len(rows)
    for row in rows:
        assert hashlib.sha256((folder/'dataset'/row['file']).read_bytes()).hexdigest()==row['sha256']
    doc=Document(folder/'IS794_Laporan_TomatoVision.docx')
    text='\n'.join(p.text for p in doc.paragraphs)
    required=['Background','Literature Study','Data Acquisition','Methodology','Results','Evaluation','Analysis and Discussion','References']
    assert all(t in text for t in required)
    assert '66,7%' in text
    assert f'{len(rows)} gambar' in text
    assert len(doc.inline_shapes)>=6
    background=text.split('Background (100–500 words)')[-1].split('Literature Study')[0]
    assert 100<=len(background.split())<=500
    decks={}
    for path in (folder/'presentasi').glob('*.pptx'):
        slides=Presentation(path).slides
        assert all(s.has_notes_slide for s in slides)
        decks[path.name]=len(slides)
        for slide in slides:
            for shape in slide.shapes:
                assert shape.left>=0 and shape.top>=0
                assert shape.left+shape.width<=Presentation(path).slide_width+10
                assert shape.top+shape.height<=Presentation(path).slide_height+10
    notebook=nbformat.read(ROOT/'notebooks/IS794_TomatoVision_Project.ipynb',as_version=4)
    nbformat.validate(notebook)
    code=[c for c in notebook.cells if c.cell_type=='code']
    assert all(c.execution_count is not None for c in code)
    assert not any(o.output_type=='error' for c in code for o in c.outputs)
    result={'images':len(rows),'report_docx_valid':True,'report_figures':len(doc.inline_shapes),
            'background_words':len(background.split()),'presentation_slides':decks,
            'executed_notebook_code_cells':len(code),
            'all_assignment_requirements_met':False,
            'remaining':['Primary data quantity and acquisition evidence are not verified.',
                         'Word-only report is requested; official assignment asks for PDF.',
                         'Team identity and unique-title check are handled in a separate user document.']}
    (folder/'audit/deliverable_validation.json').write_text(json.dumps(result,indent=2)+'\n')
    print(json.dumps(result,indent=2))

if __name__=='__main__': main()
